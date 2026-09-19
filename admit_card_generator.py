from pathlib import Path
from io import BytesIO
import math

from openpyxl import load_workbook
from PIL import Image
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas

from config import (
    SCHOOL_NAME,
    SCHOOL_SUBTITLE,
    SCHOOL_ADDRESS,
    EXAM_NAME,
    ISSUE_DATE,
    LOGO_PATH,
    OUTPUT_DIR,
)
from placeholders import create_gender_placeholders


PAGE_W, PAGE_H = A4

MARGIN_X = 22
MARGIN_TOP = 18
MARGIN_BOTTOM = 18
CARD_GAP = 22

CARD_W = PAGE_W - (MARGIN_X * 2)
CARD_H = (PAGE_H - MARGIN_TOP - MARGIN_BOTTOM - CARD_GAP) / 2


def normalize(value):
    if value is None:
        return ""
    return str(value).strip()


def normalize_key(value):
    return "".join(ch.lower() for ch in normalize(value) if ch.isalnum())


def find_value(row, *possible_names):
    normalized = {normalize_key(k): v for k, v in row.items()}
    for name in possible_names:
        key = normalize_key(name)
        if key in normalized:
            return normalize(normalized[key])
    return ""


def read_students(excel_path):
    wb = load_workbook(excel_path, data_only=True)
    ws = wb.active

    headers = [cell.value for cell in ws[1]]
    students = []

    for row in ws.iter_rows(min_row=2, values_only=True):
        if not any(v is not None and str(v).strip() for v in row):
            continue

        record = dict(zip(headers, row))
        students.append({
            "name": find_value(record, "Name of Student", "Student Name", "Name", "StudentName"),
            "class": find_value(record, "Class", "Class Name", "Grade"),
            "roll_no": find_value(record, "Roll No", "Roll No.", "Roll Number", "RollNumber"),
            "father_name": find_value(record, "Father's Name", "Father Name", "FatherName"),
            "mother_name": find_value(record, "Mother's Name", "Mother Name", "MotherName"),
            "contact": find_value(record, "Contact No", "Contact No.", "Contact", "Phone", "Mobile"),
            "gender": find_value(record, "Gender", "Sex"),
        })

    return students


def fit_text(c, text, max_width, font_name="Helvetica-Bold", start_size=11, min_size=7):
    size = start_size
    while size > min_size:
        if c.stringWidth(text, font_name, size) <= max_width:
            return size
        size -= 0.5
    return min_size


def draw_image_contain(c, image_path, x, y, w, h):
    if not image_path.exists():
        return

    try:
        img = Image.open(image_path)
        img_w, img_h = img.size
        scale = min(w / img_w, h / img_h)
        draw_w = img_w * scale
        draw_h = img_h * scale
        draw_x = x + (w - draw_w) / 2
        draw_y = y + (h - draw_h) / 2
        c.drawImage(
            ImageReader(img),
            draw_x,
            draw_y,
            width=draw_w,
            height=draw_h,
            preserveAspectRatio=True,
            mask="auto",
        )
    except Exception:
        pass


def draw_photo(c, photo_path, x, y, w=68, h=80):
    c.setStrokeColor(colors.black)
    c.setLineWidth(0.8)
    c.rect(x, y, w, h, stroke=1, fill=0)

    if photo_path.exists():
        draw_image_contain(c, photo_path, x + 2, y + 2, w - 4, h - 4)


def draw_label_value(c, label, value, x, y, label_width=70, value_width=185):
    c.setFont("Helvetica-Bold", 8.5)
    c.drawString(x, y, label)

    c.setFont("Helvetica", 8.5)
    value = value or ""
    max_width = value_width
    size = fit_text(c, value, max_width, "Helvetica", 12, 6.5)
    c.setFont("Helvetica", size)
    value_x = x + label_width
    c.drawString(value_x, y, value)

    if value:
        value_width = c.stringWidth(value, "Helvetica", size)
        c.setLineWidth(0.5)
        c.line(value_x, y - 2, value_x + value_width, y - 2)


def draw_card(c, student, x, y, w, h, placeholders_dir):
    # Outer border
    c.setStrokeColor(colors.black)
    c.setLineWidth(1.5)
    c.rect(x, y, w, h, stroke=1, fill=0)

    padding = 10
    inner_x = x + padding
    inner_top = y + h - padding

    # Logo
    logo_w = 60
    logo_h = 48
    draw_image_contain(
        c,
        LOGO_PATH,
        inner_x,
        inner_top - logo_h,
        logo_w,
        logo_h,
    )

    # Header
    header_center = x + w / 2
    c.setFillColor(colors.black)

    school_size = fit_text(c, SCHOOL_NAME, w - 105, "Helvetica-Bold", 15, 10)
    c.setFont("Helvetica-Bold", school_size)
    c.drawCentredString(header_center, inner_top - 9, SCHOOL_NAME)

    c.setFont("Helvetica-Bold", 12)
    c.drawCentredString(header_center, inner_top - 21, SCHOOL_SUBTITLE)
    c.drawCentredString(header_center, inner_top - 32, SCHOOL_ADDRESS)

    # Admit card title box
    title_w = 150
    title_h = 25
    title_x = header_center - title_w / 2
    title_y = inner_top - 66
    c.roundRect(title_x, title_y, title_w, title_h, 4, stroke=1, fill=0)
    c.setFont("Helvetica-Bold", 13)
    c.drawCentredString(header_center, title_y + 7, "ADMIT CARD")

    # Student information area
    info_top = title_y - 13
    photo_w = 68
    photo_h = 80
    photo_x = x + w - padding - photo_w
    photo_y = info_top - photo_h

    gender = student["gender"].lower()
    if gender.startswith("f") or "female" in gender:
        photo_path = placeholders_dir / "female_placeholder.png"
    else:
        photo_path = placeholders_dir / "male_placeholder.png"

    draw_photo(c, photo_path, photo_x, photo_y, photo_w, photo_h)

    text_x = inner_x
    col_gap = 12
    left_width = photo_x - text_x - col_gap

    row_y = info_top - 12
    draw_label_value(c, "Name of Student:", student["name"], text_x, row_y,
                     label_width=78, value_width=left_width - 78)

    row_y -= 20
    draw_label_value(c, "Class:", student["class"], text_x, row_y,
                     label_width=78, value_width=left_width - 78)

    row_y -= 20
    draw_label_value(c, "Roll No.:", student["roll_no"], text_x, row_y,
                     label_width=78, value_width=left_width - 78)

    row_y -= 20
    draw_label_value(c, "Father's Name:", student["father_name"], text_x, row_y,
                     label_width=78, value_width=left_width - 78)

    row_y -= 20
    draw_label_value(c, "Mother's Name:", student["mother_name"], text_x, row_y,
                     label_width=78, value_width=left_width - 78)

    # Contact on the right side of the information area
    contact_x = text_x + left_width * 0.53
    contact_y = row_y
    if student["contact"]:
        draw_label_value(
            c, "Contact No.:", student["contact"],
            contact_x, contact_y,
            label_width=62,
            value_width=photo_x - contact_x - 5
        )

    # Divider
    divider_y = row_y - 10
    c.setLineWidth(0.6)
    c.line(inner_x, divider_y, x + w - padding, divider_y)

    # Examination box
    exam_y = divider_y - 57
    c.roundRect(inner_x, exam_y, w - 2 * padding, 50, 3, stroke=1, fill=0)

    c.setFont("Helvetica-Bold", 12)
    c.drawCentredString(header_center, exam_y + 37, EXAM_NAME)

    c.setLineWidth(0.4)
    c.line(inner_x, exam_y + 29, x + w - padding, exam_y + 29)

    c.setFont("Helvetica-Bold", 8.5)
    c.drawString(inner_x + 8, exam_y + 19, "Examination:")
    c.setFont("Helvetica", 10)
    c.drawString(inner_x + 72, exam_y + 19, EXAM_NAME)

    c.setFont("Helvetica-Bold", 8.5)
    c.drawString(inner_x + 8, exam_y + 5, "Issue Date:")
    c.setFont("Helvetica", 10)
    c.drawString(inner_x + 72, exam_y + 5, ISSUE_DATE)

    # Note
    note_y = exam_y - 14
    c.setFont("Helvetica-Bold", 8)
    c.drawString(inner_x, note_y, "Note:")
    c.setFont("Helvetica", 7.5)
    c.drawString(
        inner_x + 27,
        note_y,
        "Keep this card safely and must bring to the exam venue on every exam date."
    )

    # Signature row
    sig_y = y + 21
    sig_w = 75
    positions = [
        (inner_x, "Principal"),
        (header_center - sig_w / 2, "School Seal"),
        (x + w - padding - sig_w, "Exam Controller"),
    ]

    for sx, label in positions:
        c.line(sx, sig_y + 12, sx + sig_w, sig_y + 12)
        c.setFont("Helvetica-Bold", 7.5)
        c.drawCentredString(sx + sig_w / 2, sig_y, label)


def generate_pdf(excel_path, output_pdf=None):
    excel_path = Path(excel_path)

    create_gender_placeholders(excel_path.parent / "assets" if False else Path("assets"))

    students = read_students(excel_path)
    if not students:
        raise ValueError("No student records were found in the Excel file.")

    placeholders_dir = Path("assets")

    if output_pdf is None:
        output_pdf = OUTPUT_DIR / "admit_cards.pdf"
    else:
        output_pdf = Path(output_pdf)

    output_pdf.parent.mkdir(parents=True, exist_ok=True)

    c = canvas.Canvas(str(output_pdf), pagesize=A4)

    for index, student in enumerate(students):
        position = index % 2

        if position == 0:
            card_y = PAGE_H - MARGIN_TOP - CARD_H
        else:
            card_y = MARGIN_BOTTOM

        draw_card(
            c,
            student,
            MARGIN_X,
            card_y,
            CARD_W,
            CARD_H,
            placeholders_dir,
        )

        if position == 1 or index == len(students) - 1:
            c.showPage()

    c.save()
    return output_pdf


if __name__ == "__main__":
    import argparse

    parser = argparse.ArgumentParser(description="Generate AMFS examination admit cards.")
    parser.add_argument("excel", help="Path to the student Excel file")
    parser.add_argument(
        "-o",
        "--output",
        default=str(OUTPUT_DIR / "admit_cards.pdf"),
        help="Output PDF path",
    )
    args = parser.parse_args()

    pdf = generate_pdf(args.excel, args.output)
    print(f"Generated: {pdf}")
